import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins: 1 inch
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Primary colors
    COLOR_PRIMARY = RGBColor(27, 67, 50)     # #1B4332 - Deep Forest Green
    COLOR_SECONDARY = RGBColor(45, 106, 79)  # #2D6A4F - Medium Forest Green
    COLOR_DARK = RGBColor(33, 37, 41)        # #212529 - Charcoal Text
    COLOR_MUTED = RGBColor(108, 117, 125)    # #6C757D - Muted Gray

    # ------------------ COVER PAGE ------------------
    cover_p1 = doc.add_paragraph()
    cover_p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_p1.paragraph_format.space_before = Pt(36)
    cover_p1.paragraph_format.space_after = Pt(8)
    run_badge = cover_p1.add_run("ACADEMIC PROJECT REPORT")
    run_badge.font.name = "Arial"
    run_badge.font.size = Pt(11)
    run_badge.font.bold = True
    run_badge.font.color.rgb = COLOR_SECONDARY

    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_title.paragraph_format.space_before = Pt(12)
    cover_title.paragraph_format.space_after = Pt(12)
    run_title = cover_title.add_run("HarmonyHub")
    run_title.font.name = "Georgia"
    run_title.font.size = Pt(36)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    cover_sub = doc.add_paragraph()
    cover_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_sub.paragraph_format.space_after = Pt(28)
    run_sub = cover_sub.add_run("A Digital Learning Platform for Universal Human Values (UHV) & Professional Ethics")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(15)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    tagline_p = doc.add_paragraph()
    tagline_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tagline_p.paragraph_format.space_after = Pt(40)
    tagline_run = tagline_p.add_run('"Understand Yourself. Build Better Relationships. Live in Harmony."')
    tagline_run.font.name = "Georgia"
    tagline_run.font.size = Pt(13)
    tagline_run.font.bold = True
    tagline_run.font.color.rgb = COLOR_PRIMARY

    # Metadata Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Course / Subject:", "Universal Human Values & Professional Ethics (UHV)"),
        ("Project Type:", "Full-Stack Web Educational Application"),
        ("Technology Stack:", "React 19, TypeScript, Vite, Tailwind CSS v4, Supabase"),
        ("Status:", "Phase 1 Completed (Responsive UI, Interactive Modules, Dashboard)"),
        ("Academic Year:", "2025 – 2026")
    ]
    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0 = row.cells[0]
        c1 = row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(4)
        p0.paragraph_format.space_before = Pt(4)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(4)
        p1.paragraph_format.space_before = Pt(4)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(10)
        r1.font.color.rgb = COLOR_DARK

        set_cell_background(c0, "F4F9F5")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, 100, 100, 120, 120)
        set_cell_margins(c1, 100, 100, 120, 120)

    doc.add_page_break()

    # ------------------ HELPER FUNCTIONS FOR SECTIONS ------------------
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Georgia"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix + " ")
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_PRIMARY
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_DARK
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix + ": ")
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_PRIMARY
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.font.color.rgb = COLOR_DARK
        return p

    # ------------------ 1. EXECUTIVE SUMMARY ------------------
    add_h1("1. Executive Summary")
    add_p("Universal Human Values (UHV) and Professional Ethics forms a cornerstone in contemporary higher technical education, emphasizing the holistic development of students beyond pure technical competencies. HarmonyHub is an interactive, digital educational web application engineered specifically to make UHV concepts accessible, engaging, and practically applicable.")
    add_p("The platform guides students through self-exploration, emotional intelligence, interpersonal harmony, ecological coexistence, and professional integrity. Designed with a calming natural color palette (forest green, warm cream, clean whites) and a contemporary glassmorphism aesthetic, HarmonyHub provides structured modules, micro-learning topics, reflective exercises, and personalized progress tracking.")

    # ------------------ 2. OBJECTIVES & CORE MOTIVATION ------------------
    add_h1("2. Project Objectives & Motivation")
    add_bullet("Facilitate Continuous Self-Exploration", "Self-Awareness")
    add_bullet("Strengthen Interpersonal Relationships and Trust", "Social Harmony")
    add_bullet("Foster Environmental Consciousness and Mutual Fulfillment with Nature", "Ecological Balance")
    add_bullet("Promote Ethical Competence in Professional Careers", "Professional Ethics")
    add_bullet("Provide Modern, Accessible Digital Learning Tools", "Pedagogical Modernization")

    # ------------------ 3. ARCHITECTURE & TECH STACK ------------------
    add_h1("3. System Architecture & Technology Stack")
    add_p("HarmonyHub was built using modern front-end engineering standards to guarantee blazingly fast load times, rock-solid type safety, and seamless responsive design across all devices.")

    # Tech Stack Table
    t_stack = doc.add_table(rows=7, cols=3)
    t_stack.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Layer / Domain", "Technology Selected", "Purpose & Architectural Role"]
    for j, h in enumerate(headers):
        cell = t_stack.rows[0].cells[j]
        set_cell_background(cell, "1B4332")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(10)
        set_cell_margins(cell, 120, 120, 150, 150)

    stack_rows = [
        ("Core UI Framework", "React 19 + TypeScript", "Component-driven declarative architecture with strict static typing preventing runtime bugs."),
        ("Build Tooling", "Vite 6.x", "Instant Hot Module Replacement (HMR) and optimized Rollup-based production bundling."),
        ("Styling System", "Tailwind CSS v4 + @tailwindcss/vite", "Modern utility-first CSS engine with custom @theme tokens for serene earth-toned aesthetics."),
        ("Navigation & Routing", "React Router v7", "Single Page Application (SPA) routing with dynamic URL parameters for modules and subpages."),
        ("Iconography", "Lucide React", "Lightweight, visually consistent vector icons for intuitive user navigation."),
        ("Backend & Auth (Ready)", "Supabase Client (@supabase/supabase-js)", "Cloud PostgreSQL database, Row Level Security, and JWT user authentication pipeline.")
    ]

    for i, row_data in enumerate(stack_rows):
        row = t_stack.rows[i+1]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            if i % 2 == 0:
                set_cell_background(cell, "F7FAF8")
            else:
                set_cell_background(cell, "FFFFFF")
            set_cell_margins(cell, 90, 90, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = COLOR_DARK
            if j == 0:
                r.font.bold = True

    # ------------------ 4. CURRICULUM & MODULE BREAKDOWN ------------------
    add_h1("4. UHV Curriculum & Module Breakdown")
    add_p("The platform is structured strictly around the four universally recognized pillars of Harmony in human life, providing 32 distinct interactive learning topics:")

    modules_info = [
        ("Module 1: Harmony in the Human Being", "8 Topics | Self-Exploration, Desire vs Need, Body-Self Synergy", [
            "Introduction to Self-Exploration: Natural Acceptance and Experiential Validation",
            "Continuous Happiness and Prosperity: The Basic Human Aspirations",
            "The Human Being as Co-existence of the Self ('I') and the Body",
            "Needs of the Self ('I') and 'Body' - Sukh and Suvidha",
            "Activities of the Self: Imaging, Analyzing, Selecting/Desiring",
            "Harmony in the Self: Realization and Understanding",
            "Harmony of the Self with the Body: Sanyama and Swasthya",
            "Programs to Ensure Sanyama and Health in Daily Life"
        ]),
        ("Module 2: Harmony in Family and Society", "8 Topics | Trust, Respect, Justice, Undivided Society", [
            "Harmony in the Family: The Basic Unit of Human Interaction",
            "Trust (Vishwas): The Foundational Value in Relationships",
            "Respect (Samman): Distinction vs Differentiation",
            "Affection, Care, and Guidance: Nurturing Mutual Trust",
            "Reverence, Glory, Gratitude, and Love: Complete Value Fulfillment",
            "Harmony in the Society: From Family Order to World Family",
            "Comprehensive Human Goals: Right Understanding, Prosperity, Fearlessness, Co-existence",
            "Universal Human Order (Sarvabhauma Vyavastha)"
        ]),
        ("Module 3: Harmony in Nature and Existence", "8 Topics | Four Orders of Nature, Cyclic Processes, Co-existence", [
            "Understanding Harmony in Nature: The Four Orders",
            "Physical Order (Padartha Avastha): Soil, Water, Metals, and Minerals",
            "Bio Order (Prana Avastha): Plants, Trees, and Microorganisms",
            "Animal Order (Jeeva Avastha): Animals and Birds",
            "Human Order (Gyana Avastha): Self-Conscious Entities",
            "Interconnectedness and Mutual Fulfillment (Parasparta)",
            "Recyclability and Self-Regulation in Nature",
            "Existence as Co-existence (Sah-astitva) in Space"
        ]),
        ("Module 4: Professional Ethics", "8 Topics | Ethical Competence, Holistic Technologies, Sustainable Living", [
            "Definitiveness of Ethical Human Conduct",
            "Basis for Humanistic Education, Constitution, and Universal Order",
            "Competence in Professional Ethics: Clarity, Confidence, and Integrity",
            "Issues in Professional Ethics: Exploitation, Corruption, and Unethical Practices",
            "Holistic Technologies and Production Systems: Eco-Friendly Criteria",
            "Management Models for Mutual Fulfillment and Long-term Sustainability",
            "Case Studies of Typical Ethical Dilemmas in Engineering and Management",
            "Vision for a Holistic and Sustainable Alternative Way of Living"
        ])
    ]

    for title, subtitle, topics in modules_info:
        add_h2(title)
        add_p(subtitle, "Module Scope:")
        for top in topics:
            add_bullet(top)

    # ------------------ 5. IMPLEMENTED FEATURES & PAGES ------------------
    add_h1("5. Implemented Features & Application Pages")
    add_p("The initial release (Phase 1) delivers a complete, highly polished frontend user journey with comprehensive educational content, dynamic statistics, interactive navigation, and full responsiveness.")

    features = [
        ("Landing Page (Home)", "Features an attention-grabbing hero section, student impact metrics, 4 module overview cards with smooth hover micro-animations, core values highlights, student testimonials, and call-to-action sections."),
        ("About HarmonyHub Page", "Presents the academic mission, pedagogical vision, the AICTE/UGC UHV framework context, teaching methodology, and project stewardship values."),
        ("Student Learning Dashboard", "A personalized hub showcasing enrolled courses, real-time progress calculations, overall completion ring (SVG circular indicator), active streak counter, and recent learning activity feed."),
        ("Module Detail View", "Deep dive page for each UHV pillar featuring detailed curriculum listings, estimated completion times, difficulty levels, prerequisite indicators, and completion checkoffs."),
        ("Authentication Gateway", "Dedicated Login and Registration interfaces featuring responsive form controls, password toggles, and seamless integration hooks for Supabase Auth.")
    ]

    for name, desc in features:
        add_p(desc, bold_prefix=name)

    # ------------------ SCREENSHOTS SECTION ------------------
    add_h1("6. Visual Walkthrough & Interface Screenshots")
    add_p("Below are verified high-resolution interface captures demonstrating the aesthetic design and functional layouts created for HarmonyHub:")

    # Screenshots from artifact folder
    artifact_dir = r"C:\Users\Dell\.gemini\antigravity-ide\brain\7e0d3688-22e6-4561-a2ac-9582b045fe23"
    shots = [
        ("home_page_1792967385801.png", "Figure 1: HarmonyHub Landing Page with Hero & Module Cards"),
        ("about_page_1792967402886.png", "Figure 2: About Page - Mission, Vision & Curriculum Pillars"),
        ("dashboard_page_1792967430909.png", "Figure 3: Student Learning Dashboard with Interactive Progress Tracking"),
        ("module_page_1792967456078.png", "Figure 4: Module Detail View - Human Being & 8 Detailed Topics"),
        ("login_page_1792967446685.png", "Figure 5: Student Authentication Portal with Clean Security UI")
    ]

    for img_name, caption in shots:
        img_path = os.path.join(artifact_dir, img_name)
        if os.path.exists(img_path):
            try:
                doc.add_paragraph()
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(4)
                p_img.paragraph_format.keep_with_next = True
                run_img = p_img.add_run()
                run_img.add_picture(img_path, width=Inches(5.8))

                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(14)
                r_cap = p_cap.add_run(caption)
                r_cap.font.name = "Arial"
                r_cap.font.size = Pt(9)
                r_cap.font.italic = True
                r_cap.font.color.rgb = COLOR_MUTED
            except Exception as e:
                add_p(f"[Image: {caption} - path: {img_path}]")

    # ------------------ 7. PROJECT DIRECTORY STRUCTURE ------------------
    add_h1("7. Repository Organization & File Structure")
    add_p("The project codebase maintains clean separation of concerns, modular components, and strict TypeScript types:")

    code_p = doc.add_paragraph()
    code_p.paragraph_format.space_after = Pt(10)
    code_p.paragraph_format.line_spacing = 1.05
    dir_tree = (
        "HarmonyHub/\n"
        "├── index.html                   # HTML entry point with Google Fonts (Inter, Outfit)\n"
        "├── package.json                 # Project dependencies and npm scripts\n"
        "├── vite.config.ts               # Vite configuration with @tailwindcss/vite plugin\n"
        "├── tsconfig.json                # TypeScript compiler configuration\n"
        "├── .env.example                 # Template for Supabase URL and Anon Key\n"
        "├── src/\n"
        "│   ├── main.tsx                 # React application mounting point\n"
        "│   ├── App.tsx                  # BrowserRouter and route mapping\n"
        "│   ├── index.css                # Global Tailwind v4 styles, custom tokens & utilities\n"
        "│   ├── types/\n"
        "│   │   └── index.ts             # TypeScript interfaces (Module, Topic, User)\n"
        "│   ├── data/\n"
        "│   │   └── modules.ts           # 4 UHV Modules and 32 learning topics dataset\n"
        "│   ├── lib/\n"
        "│   │   └── supabase.ts          # Supabase client initialization & connection status\n"
        "│   ├── components/\n"
        "│   │   ├── layout/\n"
        "│   │   │   ├── Navbar.tsx       # Glassmorphism header with responsive mobile drawer\n"
        "│   │   │   └── Footer.tsx       # Dark green footer with quick links and copyright\n"
        "│   │   └── ui/\n"
        "│   │       └── ModuleCard.tsx   # Color-themed card component with progress badges\n"
        "│   └── pages/\n"
        "│       ├── LandingPage.tsx      # Landing hero, stats, cards, testimonials\n"
        "│       ├── AboutPage.tsx        # Mission, pedagogical approach, UHV values\n"
        "│       ├── DashboardPage.tsx    # Enrolled courses, circular progress, streak\n"
        "│       ├── ModuleDetailPage.tsx # Topic breakdown, prerequisites, start button\n"
        "│       ├── LoginPage.tsx        # Student login form with security validation\n"
        "│       └── RegisterPage.tsx     # Student registration form\n"
    )
    r_code = code_p.add_run(dir_tree)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = COLOR_DARK

    # ------------------ 8. DATABASE & BACKEND ARCHITECTURE ------------------
    add_h1("8. Supabase Database Schema & Security Design")
    add_p("HarmonyHub is architected for turnkey integration with Supabase (PostgreSQL with Row Level Security). Below is the recommended database schema for user profiles, module tracking, and reflective journal submissions:")

    sql_p = doc.add_paragraph()
    sql_p.paragraph_format.space_after = Pt(10)
    sql_code = (
        "-- 1. User Profiles linked to Supabase Auth\n"
        "CREATE TABLE public.profiles (\n"
        "    id UUID REFERENCES auth.users ON DELETE CASCADE PRIMARY KEY,\n"
        "    full_name TEXT NOT NULL,\n"
        "    email TEXT UNIQUE NOT NULL,\n"
        "    role TEXT DEFAULT 'student',\n"
        "    avatar_url TEXT,\n"
        "    created_at TIMESTAMPTZ DEFAULT NOW()\n"
        ");\n\n"
        "-- 2. Student Topic Progress Tracking\n"
        "CREATE TABLE public.user_progress (\n"
        "    id UUID DEFAULT gen_random_uuid() PRIMARY KEY,\n"
        "    user_id UUID REFERENCES public.profiles(id) ON DELETE CASCADE,\n"
        "    module_id TEXT NOT NULL,\n"
        "    topic_id TEXT NOT NULL,\n"
        "    completed BOOLEAN DEFAULT FALSE,\n"
        "    completed_at TIMESTAMPTZ,\n"
        "    UNIQUE(user_id, topic_id)\n"
        ");\n\n"
        "-- 3. Self-Exploration Reflective Journal\n"
        "CREATE TABLE public.reflections (\n"
        "    id UUID DEFAULT gen_random_uuid() PRIMARY KEY,\n"
        "    user_id UUID REFERENCES public.profiles(id) ON DELETE CASCADE,\n"
        "    topic_id TEXT NOT NULL,\n"
        "    reflection_text TEXT NOT NULL,\n"
        "    created_at TIMESTAMPTZ DEFAULT NOW()\n"
        ");"
    )
    r_sql = sql_p.add_run(sql_code)
    r_sql.font.name = "Consolas"
    r_sql.font.size = Pt(8.5)
    r_sql.font.color.rgb = COLOR_DARK

    # ------------------ 9. SETUP & DEPLOYMENT GUIDE ------------------
    add_h1("9. Setup & Local Execution Guide")
    add_bullet("git clone <repository_url> && cd HarmonyHub", "Step 1: Clone Repository")
    add_bullet("npm install", "Step 2: Install Node Dependencies")
    add_bullet("cp .env.example .env.local (configure Supabase URL & Key)", "Step 3: Setup Environment Variables")
    add_bullet("npm run dev (launches Vite HMR on http://localhost:5173)", "Step 4: Launch Development Server")
    add_bullet("npm run build (verifies strict TypeScript compilation and asset bundling)", "Step 5: Production Build")

    # ------------------ 10. CONCLUSION & FUTURE SCOPE ------------------
    add_h1("10. Conclusion & Roadmap")
    add_p("HarmonyHub successfully demonstrates how contemporary web engineering and human-centered design can elevate ethical education. By transforming passive theoretical reading into an engaging, structured, and reflective digital journey, students cultivate genuine self-awareness, social responsibility, and ethical leadership.")
    add_p("Future enhancements scheduled for Phase 2 include live Supabase authentication, interactive topic reading panes with self-assessment quizzes, AI-assisted reflective journal prompts, and faculty grading dashboards.")

    # Save document
    out_path = r"e:\megaproject\HarmonyHub\HarmonyHub_Project_Report.docx"
    doc.save(out_path)
    print(f"SUCCESS: Report saved to {out_path}")

if __name__ == "__main__":
    create_report()
